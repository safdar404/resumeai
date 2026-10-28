"""Re-open generated files and verify their content, structure and layout."""
from __future__ import annotations

import io
import re

from pydantic import BaseModel, Field

from src.models.profile import CandidateProfile
from src.models.resume import ResumeDocument
from src.services.dates import total_experience_years
from src.services.text_utils import numbers_in


class Check(BaseModel):
    name: str
    passed: bool
    detail: str = ""
    severity: str = "error"     # error | warning | info


class FileReport(BaseModel):
    kind: str
    checks: list[Check] = Field(default_factory=list)
    pages: int | None = None
    text_extractable: bool | None = None
    reading_order: bool | None = None
    blank_pages: list[int] = Field(default_factory=list)
    clipped: bool = False
    sparse_last_page: bool = False
    missing_text: list[str] = Field(default_factory=list)

    @property
    def passed(self) -> bool:
        return all(c.passed for c in self.checks if c.severity == "error")


def _norm(t: str) -> str:
    t = t.replace("\u00ad", "").replace("’", "'").replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", t).strip().lower()


def expected_snippets(doc: ResumeDocument) -> list[str]:
    out = [doc.name] + list(doc.contact_items)
    for s in doc.enabled_sections():
        out.append(s.title)
        if s.text:
            out.append(s.text[:60])
        for e in s.entries:
            out += [x for x in (e.heading, e.subheading, e.dates) if x]
            out += [b[:50] for b in e.bullets]
        out += [i[:50] for i in s.items]
    return [o for o in out if o and o.strip()]


def _find_missing(text: str, snippets: list[str]) -> list[str]:
    nt = _norm(text)
    nt_nospace = nt.replace(" ", "")
    missing = []
    for s in snippets:
        ns = _norm(s)
        if ns in nt or ns.replace(" ", "") in nt_nospace:
            continue
        missing.append(s)
    return missing


def _allowed_numbers(p: CandidateProfile, doc: ResumeDocument) -> set[str]:
    allowed = numbers_in(p.corpus()) | numbers_in(" ".join(doc.contact_items))
    for s in doc.sections:
        for e in s.entries:
            allowed |= numbers_in(e.dates)
    y = total_experience_years(p.experience)
    allowed.add(str(int(y)))
    allowed |= {str(i) for i in range(1, 6)}  # page numbers / list indices
    return allowed


def validate_pdf(pdf: bytes, doc: ResumeDocument, p: CandidateProfile, max_pages: int | None,
                 expected_order: list[str]) -> FileReport:
    import pymupdf as fitz

    rep = FileReport(kind="pdf")
    try:
        d = fitz.open(stream=pdf, filetype="pdf")
    except Exception as exc:
        rep.checks.append(Check(name="File opens", passed=False, detail=type(exc).__name__))
        return rep
    rep.pages = len(d)
    rep.checks.append(Check(name="File opens", passed=True, detail=f"{len(d)} page(s)"))
    texts, clipped = [], []
    for i, page in enumerate(d):
        t = page.get_text("text")
        texts.append(t)
        if not t.strip():
            rep.blank_pages.append(i + 1)
        rect = page.rect
        for b in page.get_text("dict")["blocks"]:
            for line in b.get("lines", []):
                for span in line["spans"]:
                    x0, y0, x1, y1 = span["bbox"]
                    if span["text"].strip() and (x0 < rect.x0 - 0.5 or y0 < rect.y0 - 0.5 or
                                                 x1 > rect.x1 + 0.5 or y1 > rect.y1 + 0.5):
                        clipped.append(span["text"][:30])
    full = "\n".join(texts)
    rep.text_extractable = len(full.strip()) > 50
    rep.checks.append(Check(name="Text is selectable and extractable", passed=bool(rep.text_extractable),
                            detail=f"{len(full.split())} words extracted"))
    rep.checks.append(Check(name="No blank pages", passed=not rep.blank_pages,
                            detail=f"blank: {rep.blank_pages}" if rep.blank_pages else ""))
    rep.clipped = bool(clipped)
    rep.checks.append(Check(name="No text beyond page edges", passed=not clipped,
                            detail=", ".join(clipped[:3])))
    if len(texts) > 1:
        last_lines = [l for l in texts[-1].splitlines() if l.strip()]
        rep.sparse_last_page = len(last_lines) < 5
        rep.checks.append(Check(name="Last page is not nearly empty", passed=not rep.sparse_last_page, severity="warning",
                                detail=f"{len(last_lines)} line(s) on the last page"))
    if max_pages:
        rep.checks.append(Check(name="Page count matches target", passed=len(d) <= max_pages, severity="warning",
                                detail=f"{len(d)} of {max_pages} allowed"))
    rep.missing_text = _find_missing(full, expected_snippets(doc))
    rep.checks.append(Check(name="Expected sections and essential text present", passed=not rep.missing_text,
                            detail=("missing: " + "; ".join(m[:40] for m in rep.missing_text[:5])) if rep.missing_text else
                            "all sections, entries and bullets found"))
    nf = _norm(full)
    ok_identity = _norm(p.full_name) in nf and (not p.contact.email or _norm(p.contact.email) in nf.replace(" ", "") or _norm(p.contact.email) in nf)
    rep.checks.append(Check(name="Name and contact details intact", passed=ok_identity))
    # reading order
    ordered, last = True, -1
    for item in expected_order:
        idx = nf.find(_norm(item), max(last, 0))
        if idx < 0:
            idx = nf.find(_norm(item))
            if idx >= 0 and idx < last:
                ordered = False
        if idx >= 0:
            last = max(last, idx)
    rep.reading_order = ordered
    rep.checks.append(Check(name="Sensible reading order", passed=ordered, severity="warning",
                            detail="headings extracted in visual order" if ordered else "some headings extract out of order"))
    extra = numbers_in(full) - _allowed_numbers(p, doc)
    rep.checks.append(Check(name="No unsupported numbers introduced", passed=not extra,
                            detail=", ".join(sorted(extra)[:6])))
    return rep


def docx_text(data: bytes) -> tuple[str, bool, list[str]]:
    from docx import Document

    d = Document(io.BytesIO(data))
    parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.extend(pp.text for pp in cell.paragraphs)
    header_text = any(pp.text.strip() for s in d.sections for part in (s.header, s.footer) for pp in part.paragraphs)
    styles = sorted({pp.style.name for pp in d.paragraphs if pp.style is not None})
    return "\n".join(parts), header_text, styles


def validate_docx(data: bytes, doc: ResumeDocument, p: CandidateProfile) -> FileReport:
    rep = FileReport(kind="docx")
    try:
        text, header_text, styles = docx_text(data)
    except Exception as exc:
        rep.checks.append(Check(name="File opens", passed=False, detail=type(exc).__name__))
        return rep
    rep.checks.append(Check(name="File opens", passed=True))
    rep.text_extractable = bool(text.strip())
    rep.missing_text = _find_missing(text, expected_snippets(doc))
    rep.checks.append(Check(name="Expected sections and essential text present", passed=not rep.missing_text,
                            detail=("missing: " + "; ".join(m[:40] for m in rep.missing_text[:5])) if rep.missing_text else
                            "all sections, entries and bullets found"))
    rep.checks.append(Check(name="No essential text in headers/footers", passed=not header_text))
    rep.checks.append(Check(name="Section headings use Word heading styles", passed="Heading 1" in styles or "Heading 2" in styles,
                            severity="warning"))
    nt = _norm(text)
    rep.checks.append(Check(name="Name and contact details intact",
                            passed=_norm(p.full_name) in nt and (not p.contact.email or _norm(p.contact.email) in nt)))
    extra = numbers_in(text) - _allowed_numbers(p, doc)
    rep.checks.append(Check(name="No unsupported numbers introduced", passed=not extra, detail=", ".join(sorted(extra)[:6])))
    rep.checks.append(Check(name="Page count", passed=True, severity="info",
                            detail="Word computes pagination on open; layout mirrors the PDF's margins and sizes."))
    return rep


def page_images(pdf: bytes, zoom: float = 1.3, max_pages: int = 4) -> list[bytes]:
    import pymupdf as fitz

    d = fitz.open(stream=pdf, filetype="pdf")
    out = []
    for i, page in enumerate(d):
        if i >= max_pages:
            break
        out.append(page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False).tobytes("png"))
    return out
