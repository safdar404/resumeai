import io

import pymupdf as fitz
import pytest
from docx import Document

from src.agents.drafting import build_document
from src.models.profile import CandidateProfile, Education, Experience
from src.rendering.docx_renderer import render_docx
from src.rendering.export_validation import validate_docx, validate_pdf
from src.rendering.pdf_renderer import render_pdf
from src.templates.registry import TEMPLATES, RenderOptions, get_template

INTL_NAMES = ["Zoë Łukasiewicz-Nguyễn", "Δημήτρης Παπαδόπουλος", "Иван Петров", "Søren Ødegård-Çelik"]


def test_fifteen_distinct_templates():
    assert len(TEMPLATES) == 15
    assert len({t.id for t in TEMPLATES}) == 15
    sigs = {(t.font, t.heading_style, t.layout, t.accent, tuple(t.section_order[:4]), t.name_align) for t in TEMPLATES}
    assert len(sigs) == 15
    assert {t.id for t in TEMPLATES if t.presentation_oriented} == {"modern_professional", "creative_portfolio"}


@pytest.mark.parametrize("spec", TEMPLATES, ids=lambda t: t.id)
def test_every_template_renders_and_validates(spec, gis_profile):
    opts = RenderOptions(page_target=spec.default_pages)
    doc = build_document(gis_profile, None, [], spec, opts)
    pdf, info = render_pdf(doc, spec, opts)
    rep = validate_pdf(pdf, doc, gis_profile, None if opts.page_target == "academic" else int(opts.page_target),
                       info.expected_order)
    assert rep.passed, [c for c in rep.checks if not c.passed]
    assert not rep.blank_pages and not rep.clipped and rep.text_extractable
    d = render_docx(doc, spec, opts, font_key=info.font)
    drep = validate_docx(d, doc, gis_profile)
    assert drep.passed, [c for c in drep.checks if not c.passed]


def test_presentation_template_has_single_column_export(gis_profile):
    spec = get_template("creative_portfolio")
    doc = build_document(gis_profile, None, [], spec, RenderOptions())
    _, info = render_pdf(doc, spec, RenderOptions())
    _, ats = render_pdf(doc, spec, RenderOptions(), force_single=True)
    assert not info.single_column and ats.single_column


def _long_profile() -> CandidateProfile:
    p = CandidateProfile(full_name="Long History Example")
    p.contact.email = "long@example.com"
    for i in range(9):
        p.experience.append(Experience(title=f"Analyst Level {i}", employer=f"Fictional Org {i}", start_date=f"{2000 + 2 * i}",
                                       end_date=f"{2001 + 2 * i}", bullets=[f"Delivered fictional workstream {j} with "
                                                                            "careful documentation and stakeholder reviews"
                                                                            for j in range(6)]))
    p.education.append(Education(degree="BSc", institution="Example University"))
    return p


def test_long_resume_overflow_is_reported_not_truncated():
    p = _long_profile()
    spec, opts = get_template("classic_ats"), RenderOptions(page_target="1")
    doc = build_document(p, None, [], spec, opts)
    pdf, info = render_pdf(doc, spec, opts)
    assert info.overflow and info.pages > 1
    assert info.body_size >= 10.0
    text = "".join(pg.get_text() for pg in fitz.open(stream=pdf, filetype="pdf"))
    for e in p.experience:                       # nothing silently deleted
        assert e.employer in text
        for b in e.bullets:
            assert b[:40] in " ".join(text.split())


def test_two_page_target_fits_long_resume():
    p = _long_profile()
    spec, opts = get_template("compact_executive"), RenderOptions(page_target="2")
    pdf, info = render_pdf(build_document(p, None, [], spec, opts), spec, opts)
    assert info.pages <= 2 and not info.overflow


def test_letter_size_and_justify(gis_profile):
    spec = get_template("finance")
    opts = RenderOptions(page_size="Letter", alignment="justify")
    pdf, _ = render_pdf(build_document(gis_profile, None, [], spec, opts), spec, opts)
    page = fitz.open(stream=pdf, filetype="pdf")[0]
    assert round(page.rect.width) == 612 and round(page.rect.height) == 792


@pytest.mark.parametrize("name", INTL_NAMES)
def test_international_names_survive_pdf_and_docx(name, gis_profile):
    gis_profile.full_name = name
    for tid in ("classic_ats", "minimal_elegant", "creative_portfolio"):
        spec, opts = get_template(tid), RenderOptions()
        doc = build_document(gis_profile, None, [], spec, opts)
        pdf, info = render_pdf(doc, spec, opts)
        assert name in fitz.open(stream=pdf, filetype="pdf")[0].get_text(), (tid, info.font)
        d = Document(io.BytesIO(render_docx(doc, spec, opts)))
        assert d.paragraphs[0].text == name


def test_pdf_and_docx_share_content(gis_profile):
    spec, opts = get_template("ai_data"), RenderOptions()
    doc = build_document(gis_profile, None, [], spec, opts)
    pdf, _ = render_pdf(doc, spec, opts)
    pdf_text = " ".join(fitz.open(stream=pdf, filetype="pdf")[0].get_text().split())
    d = Document(io.BytesIO(render_docx(doc, spec, opts)))
    docx_text = " ".join(" ".join(p.text for p in d.paragraphs).split())
    for e in gis_profile.experience:
        for b in e.bullets:
            assert b[:30] in pdf_text and b[:30] in docx_text


def test_validation_detects_missing_and_invented_text(gis_profile):
    spec, opts = get_template("classic_ats"), RenderOptions()
    doc = build_document(gis_profile, None, [], spec, opts)
    pdf, info = render_pdf(doc, spec, opts)
    tampered = doc.model_copy(deep=True)
    tampered.section("experience").entries[0].bullets.append("Bullet that was never rendered")
    rep = validate_pdf(pdf, tampered, gis_profile, 1, info.expected_order)
    assert rep.missing_text and not rep.passed
    bad = doc.model_copy(deep=True)
    bad.section("summary").text += " Saved 4,321 hours."
    pdf2, info2 = render_pdf(bad, spec, opts)
    rep2 = validate_pdf(pdf2, bad, gis_profile, 1, info2.expected_order)
    assert any(c.name.startswith("No unsupported numbers") and not c.passed for c in rep2.checks)


def test_disabled_sections_and_order(gis_profile):
    spec = get_template("classic_ats")
    opts = RenderOptions(disabled_sections=["languages"], section_order=["education", "experience"])
    doc = build_document(gis_profile, None, [], spec, opts)
    keys = [s.key for s in doc.enabled_sections()]
    assert "languages" not in keys and keys[0] == "education"


def test_sidebar_docx_uses_fixed_column_widths(gis_profile):
    """Regression: without a fixed layout, Word/LibreOffice widen the sidebar and push dates off the page."""
    from docx.oxml.ns import qn
    spec = get_template("modern_professional")
    d = Document(io.BytesIO(render_docx(build_document(gis_profile, None, [], spec, RenderOptions()), spec, RenderOptions())))
    tbl = d.tables[0]._tbl
    assert tbl.tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
    left, right = [int(c.get(qn("w:w"))) for c in tbl.tblGrid.findall(qn("w:gridCol"))]
    assert 0.25 < left / (left + right) < 0.38
